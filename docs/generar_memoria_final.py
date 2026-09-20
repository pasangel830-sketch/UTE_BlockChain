from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Memoria_Final_UTE_Blockchain.docx"
LOGO = ROOT / "frontend" / "public" / "brand" / "logo-ute-blockchain.png"

NAVY = "13253D"
GOLD = "B08D3C"
LIGHT_GOLD = "F5EEDC"
LIGHT_BLUE = "EAF0F6"
LIGHT_GRAY = "F2F4F7"
DARK_GRAY = "404A57"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    keep.set(qn("w:val"), "1" if value else "0")


def set_keep_lines(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepLines"))
    if keep is None:
        keep = OxmlElement("w:keepLines")
        p_pr.append(keep)
    keep.set(qn("w:val"), "1")


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    separate.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def set_update_fields(document):
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_bookmark(paragraph, name, bookmark_id):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def configure_styles(doc):
    styles = doc.styles

    def set_language(style, language="es-ES"):
        r_pr = style.element.get_or_add_rPr()
        lang = r_pr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            r_pr.append(lang)
        lang.set(qn("w:val"), language)

    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.12
    set_language(normal)

    for style_name, size, color in (
        ("Title", 29, NAVY),
        ("Subtitle", 13, GOLD),
        ("Heading 1", 20, NAVY),
        ("Heading 2", 14, NAVY),
        ("Heading 3", 11.5, GOLD),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14 if style_name != "Heading 1" else 0)
        style.paragraph_format.space_after = Pt(7)
        set_language(style)

    styles["Heading 1"].paragraph_format.page_break_before = True

    code = styles.add_style("Código", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.font.color.rgb = RGBColor.from_string("202833")
    code.paragraph_format.left_indent = Cm(0.45)
    code.paragraph_format.right_indent = Cm(0.25)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0

    caption = styles["Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string("687384")
    set_language(caption)

    quote = styles["Quote"]
    quote.font.name = "Aptos"
    quote.font.size = Pt(10.5)
    quote.font.italic = True
    quote.font.color.rgb = RGBColor.from_string(NAVY)
    quote.paragraph_format.left_indent = Cm(0.8)
    quote.paragraph_format.right_indent = Cm(0.8)
    quote.paragraph_format.space_before = Pt(7)
    quote.paragraph_format.space_after = Pt(9)
    set_language(quote)


def configure_page(section):
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.style = "Caption"
    r = p.add_run("Memoria final · UTE Blockchain Solutions  |  ")
    r.font.color.rgb = RGBColor.from_string("748091")
    add_field(p, "PAGE", "1")


def add_para(doc, text, bold_prefix=None, style=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    set_keep_lines(p)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        if isinstance(item, tuple):
            label, body = item
            r = p.add_run(label)
            r.bold = True
            p.add_run(body)
        else:
            p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(str(value))
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
        r.font.size = Pt(9)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx, value in enumerate(row):
            cell = cells[cidx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if ridx % 2:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(8.7)
            if cidx == 0:
                r.bold = True
        if widths:
            for cidx, width in enumerate(widths):
                cells[cidx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code(doc, code, source):
    caption = doc.add_paragraph(style="Caption")
    caption.add_run(source)
    set_keep_with_next(caption)
    p = doc.add_paragraph(style="Código")
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)
    p.add_run(code.strip())
    return p


def add_callout(doc, title, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_main_heading(doc, text, bookmark_id):
    p = doc.add_paragraph(text, style="Heading 1")
    add_bookmark(p, f"epigrafe_{bookmark_id}", bookmark_id)
    return p


doc = Document()
for section in doc.sections:
    configure_page(section)
    add_page_number(section)
configure_styles(doc)
set_update_fields(doc)

# Portada
cover = doc.sections[0]
cover.different_first_page_header_footer = True
if LOGO.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO), width=Inches(2.15))
    p.paragraph_format.space_after = Pt(34)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("MEMORIA FINAL")

p = doc.add_paragraph(style="Subtitle")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Plataforma blockchain para la gestión y trazabilidad de una Unión Temporal de Empresas")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
r = p.add_run("Hitos · pagos en custodia · incidencias privadas · estado de obra")
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = RGBColor.from_string(NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(130)
r = p.add_run("Hyperledger Fabric 2.5 · TypeScript · Google Cloud · Vercel")
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string(GOLD)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(12)
p.add_run("Septiembre de 2026").font.size = Pt(10.5)

doc.add_page_break()

# Resumen e índice
p = doc.add_paragraph("Resumen", style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
add_para(
    doc,
    "En este proyecto he diseñado e implementado una plataforma de coordinación para una UTE formada por cuatro constructoras y una Administración pública. Mi objetivo ha sido resolver un problema de confianza: todas las partes necesitan compartir hitos, incidencias y pagos, pero ninguna debería controlar por sí sola el historial común. La solución combina una red permisionada Hyperledger Fabric, cuatro contratos inteligentes en TypeScript, una API Express y una interfaz Next.js.",
)
add_para(
    doc,
    "El resultado aporta un registro común, trazable y resistente a modificaciones unilaterales. La finalización de un hito y la creación de su pago en custodia se ejecutan dentro de una misma transacción de Fabric; la autorización posterior exige el endoso de la empresa responsable y de la Administración, emite el evento PagoAutorizado y activa una integración bancaria simulada. Para equilibrar transparencia y confidencialidad, mantengo la información general en el canal y los datos sensibles de las incidencias en colecciones privadas. Los documentos y fotografías se conservan fuera de la cadena y se vinculan mediante su huella SHA-256.",
)
add_para(
    doc,
    "He desplegado la red y la API en Google Cloud, he protegido el acceso público mediante HTTPS con Caddy, he preparado el frontend para Vercel y he separado la monitorización en una segunda máquina con Prometheus y Grafana. La validación actual ejecuta 43 pruebas unitarias —una más que el recuento previo de 42—, además de pruebas de integración, verificación del flujo completo y una medición de consumo optimizada cercana a 1 GB de contenedores.",
)

p = doc.add_paragraph("Índice", style="Title")
p.paragraph_format.space_before = Pt(18)
toc = doc.add_paragraph()
add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "Pulse con el botón derecho y seleccione «Actualizar campo» si Word no lo actualiza automáticamente.")

# 1. Introducción
add_main_heading(doc, "1. Introducción", 1)
doc.add_heading("1.1. ¿Por qué Fabric y no una base de datos compartida?", level=2)
add_para(
    doc,
    "Ninguna empresa quiere que otra gestione en exclusiva una base de datos central. Esta idea fue mi punto de partida. En una UTE, las compañías colaboran durante una obra concreta, pero mantienen intereses, responsabilidades y sistemas propios. La Administración, además, necesita comprobar que los hitos que justifican pagos son auténticos y que el procedimiento ha seguido las reglas acordadas. Una base de datos tradicional puede guardar toda esa información, pero siempre deja una pregunta abierta: ¿quién administra la base, quién puede corregirla y quién demuestra que una corrección no ha alterado el pasado?",
)
add_para(
    doc,
    "Podría haber creado una aplicación convencional con PostgreSQL, permisos y un registro de auditoría. Técnicamente habría sido más sencilla y, para una organización única, probablemente sería la opción adecuada. Sin embargo, no habría resuelto el núcleo del caso: la neutralidad entre partes que no quieren delegar la verdad compartida en una de ellas. El administrador de esa base conservaría privilegios técnicos para cambiar registros, restaurar copias o decidir qué historial se considera válido. El resto tendría que confiar en sus procedimientos.",
)
add_para(
    doc,
    "Con Hyperledger Fabric sustituyo esa confianza unilateral por reglas consensuadas. Cada organización posee su identidad MSP, conserva una copia del ledger en su peer y participa en el endoso según la operación. Una transacción solo se incorpora si reúne las firmas exigidas y supera la validación del canal. Después, los bloques encadenan sus hashes, de modo que reescribir una operación antigua rompería la continuidad criptográfica y dejaría de coincidir con las copias del resto de participantes.",
)
add_callout(
    doc,
    "Idea central. ",
    "No he utilizado blockchain porque una base de datos sea incapaz de almacenar hitos, sino porque el propietario de la verdad no debía ser una empresa concreta. Fabric convierte esa condición de negocio en identidades, firmas y políticas verificables.",
    LIGHT_GOLD,
)

doc.add_heading("1.2. Problema abordado", level=2)
add_para(
    doc,
    "En una obra ejecutada por varias empresas aparecen documentos, certificaciones, incidencias y pagos que atraviesan fronteras organizativas. Si cada parte trabaja en su propio sistema, surgen versiones diferentes de un mismo hecho, duplicidades, demoras y discusiones sobre quién comunicó qué y cuándo. Si todas trabajan sobre el sistema de una sola empresa, el problema cambia de forma, pero no desaparece: las demás dependen de quien controla la infraestructura.",
)
add_para(
    doc,
    "He modelado el proyecto alrededor de cuatro necesidades concretas: registrar el avance de los hitos; retener el pago hasta la autorización pública; compartir incidencias sin revelar datos sensibles a quien no participa en el lote; y disponer de una vista agregada del estado de la obra. A estas funciones he añadido evidencias documentales, exploración de bloques y monitorización operativa.",
)

doc.add_heading("1.3. Objetivos", level=2)
add_bullets(
    doc,
    [
        ("Objetivo general. ", "Construir un prototipo funcional que permita a las empresas de una UTE y a la Administración compartir un historial verificable sin depender de un administrador único."),
        ("Trazabilidad. ", "Relacionar hitos, pagos, incidencias, identidades MSP, endosos, transacciones, bloques y hashes."),
        ("Automatización. ", "Crear el pago en custodia al completar el hito y notificar al banco simulado cuando la Administración lo autoriza."),
        ("Privacidad selectiva. ", "Publicar en el canal lo necesario para coordinarse y reservar el detalle sensible a los socios del lote mediante PDC."),
        ("Integridad documental. ", "Mantener los binarios fuera de Fabric y anclar su SHA-256 para detectar cualquier alteración."),
        ("Despliegue y operación. ", "Separar frontend, API, red y monitorización; aplicar TLS; y observar disponibilidad, latencia y errores de endoso."),
        ("Validación. ", "Cubrir las reglas de negocio con pruebas unitarias e integración y comprobar el flujo completo en local y en nube."),
    ],
)

doc.add_heading("1.4. Alcance y criterio de éxito", level=2)
add_para(
    doc,
    "He planteado el resultado como un producto mínimo viable académico, no como un sistema listo para gestionar dinero público real. Considero que el prototipo cumple su propósito cuando una constructora puede crear y completar un hito con evidencia, Fabric crea el pago asociado sin una ventana de inconsistencia, la Administración lo autoriza con la política correcta, el evento llega al banco simulado y el Explorer permite reconstruir lo ocurrido. Del mismo modo, un socio debe poder leer el detalle privado de su lote y un no socio debe recibir una denegación comprobable.",
)
add_para(
    doc,
    "El alcance no incluye una PKI de producción, un banco real, firma individual con HSM, almacenamiento documental con ciclo de vida definitivo ni alta disponibilidad geográfica. He dejado estas carencias explícitas porque ocultarlas produciría una valoración técnica menos útil que reconocer qué demuestra el prototipo y qué exigiría una implantación real.",
)

# 2. Stack tecnológico
add_main_heading(doc, "2. Stack tecnológico", 2)
doc.add_heading("2.1. Visión general", level=2)
add_para(
    doc,
    "He organizado el código como un monorepositorio con carpetas independientes para red, chaincodes, backend, frontend, monitorización y documentación. Esta separación me ha permitido evolucionar cada capa sin mezclar responsabilidades y reproducir los entornos mediante scripts y Docker Compose.",
)
add_table(
    doc,
    ["Capa", "Tecnología", "Función en el proyecto"],
    [
        ["Ledger", "Hyperledger Fabric 2.5.16", "Red permisionada, consenso, endoso, historial y world state."],
        ["Contratos", "TypeScript, Node.js 18, fabric-contract-api", "Reglas de hitos, pagos, incidencias y estado de obra."],
        ["API", "Node.js 24.20.0, Express 5.2.1", "Autenticación, control de acceso, Gateway, evidencias y banco simulado."],
        ["Gateway", "@hyperledger/fabric-gateway 1.12.0 y gRPC", "Firma, evaluación, envío y escucha de eventos/bloques."],
        ["Frontend", "Next.js 15.5.24, React 19.1, Tailwind CSS 3.4", "Interfaz de las cinco organizaciones y ocho pantallas."],
        ["Contenedores", "Docker Engine y Docker Compose", "Aislamiento y reproducción de peers, orderers, API y monitorización."],
        ["Nube", "Google Cloud Compute Engine y GCS", "Ejecución de Fabric/API, monitorización y evidencias."],
        ["Publicación web", "Vercel Hobby", "Alojamiento previsto para el frontend Next.js."],
        ["HTTPS", "Caddy 2.11.4 y Let's Encrypt", "Terminación TLS y proxy inverso hacia API y Grafana."],
        ["Observabilidad", "Prometheus 3.14.0 y Grafana 12.4.1", "Métricas, paneles y reglas de alerta."],
        ["Pruebas", "Jest y scripts shell de integración", "Validación unitaria, de red, API, PDC, UI y flujo completo."],
    ],
    [2.6, 5.3, 8.0],
)

doc.add_heading("2.2. Backend", level=2)
add_para(
    doc,
    "He desarrollado la API en TypeScript sobre Express. La autenticación usa JWT con sesiones de ocho horas y asigna cada cuenta a un MSP. La API incorpora Helmet, CORS, limitación de peticiones, Swagger, carga multipart con Multer y métricas mediante prom-client. La conexión con Fabric se realiza por gRPC con TLS y keepalive, reutilizando gateways por combinación de organización firmante y peer de entrada.",
)
add_para(
    doc,
    "La API no es solo una pasarela HTTP. También traduce errores técnicos de Fabric a mensajes comprensibles, selecciona los endosantes, calcula el estado agregado, conserva el índice de evidencias, escucha bloques para el Explorer y mantiene un listener del evento PagoAutorizado. Esta concentración de orquestación en el backend ha sido deliberada: el navegador nunca maneja claves privadas de Fabric.",
)

doc.add_heading("2.3. Frontend", level=2)
add_para(
    doc,
    "He construido el frontend con el App Router de Next.js y componentes de React. Tailwind CSS me ha permitido mantener una estética institucional común y adaptar la interfaz a escritorio y móvil. El token JWT se usa para consumir la API, mientras que la sesión decodificada permite mostrar el perfil de organización, ocultar acciones no autorizadas y contextualizar lote, oficio y participación.",
)
add_para(
    doc,
    "La interfaz incluye login, panel inicial, hitos, pagos, incidencias, estado de obra, Explorer y monitorización. No he intentado duplicar la seguridad de Fabric únicamente con botones: la UI mejora la experiencia, pero las operaciones relevantes vuelven a validarse en la API y, cuando procede, en las políticas de la red.",
)

doc.add_heading("2.4. Red, contenedores y automatización", level=2)
add_para(
    doc,
    "Docker encapsula peers, orderers, CLI, chaincode, API, Caddy, Prometheus y Grafana. He mantenido un perfil diario reducido —tres orderers, peer de Empresa A y peer de Administración— y un perfil completo con los cinco peers. Ambos usan el mismo perfil de canal UteFull; por tanto, las políticas no cambian entre desarrollo y demostración.",
)
add_para(
    doc,
    "El Makefile reúne las operaciones repetibles: generación criptográfica, arranque de red, creación del canal, despliegue de chaincode, inicialización, pruebas, seed, API, UI, PDC y monitorización. Esta automatización ha sido esencial para evitar configuraciones manuales irreproducibles.",
)

doc.add_heading("2.5. Google Cloud, Vercel y almacenamiento", level=2)
add_para(
    doc,
    "Para el despliegue real utilicé una VM e2-standard-4 denominada fabric-ute en Google Cloud, dentro de una VPC privada. En ella ejecuté cinco peers, tres orderers, la API y Caddy. Reservé una IP pública, asocié el dominio ute-tfm.duckdns.org y regeneré los certificados TLS de Fabric con los SAN necesarios. Los puertos gRPC no se publican a Internet; el acceso exterior se concentra en 80/443.",
)
add_para(
    doc,
    "Las evidencias se almacenan en un bucket de Google Cloud Storage mediante la identidad de servicio de la VM, sin incluir una clave JSON en el repositorio. En local, el mismo contrato de almacenamiento utiliza disco. El frontend se preparó como proyecto Vercel con el framework Next.js y la URL HTTPS de la API. El último ensayo documentado dejó el proyecto creado pero con respuesta 404 hasta ajustar Root Directory a frontend y ejecutar un redeploy; por ello describo Vercel como integración preparada y parcialmente validada, no como un cierre que no consta en las pruebas.",
)

doc.add_heading("2.6. Monitorización y validación", level=2)
add_para(
    doc,
    "Separé la observabilidad en una segunda VM e2-small. Prometheus recopila métricas de los cinco peers, los tres orderers y la API a través de la VPC, conserva tres días y alimenta un dashboard provisionado en Grafana. Caddy publica Grafana bajo /grafana para que la aplicación pueda mostrarlo en modo kiosco.",
)
add_para(
    doc,
    "Para validar el código uso Jest en cada chaincode y scripts de integración para red, API, PDC, UI y flujo hito-pago. En la revisión final ejecuté 13 pruebas de HitoContract, 11 de PagoContract, 13 de IncidenciaContract y 6 de EstadoObraContract: 43 en total. El número 42 correspondía a un recuento anterior y quedó superado al añadirse una prueba más en HitoContract.",
)

# 3. Plan de negocio
add_main_heading(doc, "3. Plan de negocio y elección del caso de uso", 3)
doc.add_heading("3.1. Motivo de la elección", level=2)
add_para(
    doc,
    "Elegí la gestión de una UTE porque me pareció un caso en el que blockchain aporta una respuesta más convincente que en escenarios con un único propietario. Las empresas necesitan colaborar, pero siguen siendo entidades independientes. Comparten el objetivo de terminar la obra y, al mismo tiempo, deben proteger márgenes, métodos, costes y responsabilidades. La Administración actúa como tercera parte que supervisa y autoriza, sin que resulte razonable entregarle toda la información privada de cada lote.",
)
add_para(
    doc,
    "El sector de la construcción también ofrece procesos fáciles de reconocer: un hito pasa por varios estados, una evidencia acredita su terminación, una certificación genera un pago y una incidencia puede contener información técnica o económica sensible. Esta combinación me permitió demostrar identidad, endoso, privacidad, eventos, inmutabilidad y auditoría dentro de un único relato funcional.",
)

doc.add_heading("3.2. Necesidad de mercado", level=2)
add_para(
    doc,
    "El problema comercial que planteo no es la falta de software de gestión de obras. Existen ERP, gestores documentales y plataformas BIM. La necesidad específica es disponer de una capa neutral de confianza entre empresas y Administración que no sustituya necesariamente esos sistemas, sino que certifique los hitos de coordinación más relevantes.",
)
add_para(
    doc,
    "En la revisión exploratoria que acompañó al proyecto no encontré en España una solución de uso extendido que reuniera en un mismo producto el seguimiento operativo de una obra en UTE, el pago en custodia condicionado por políticas y una red Hyperledger Fabric con privacidad por lote. No presento esta observación como prueba de inexistencia absoluta, sino como una oportunidad razonable de diferenciación: hay soluciones parciales, pero no identifiqué una oferta equivalente con este alcance conjunto.",
)

doc.add_heading("3.3. Clientes, usuarios y propuesta de valor", level=2)
add_table(
    doc,
    ["Actor", "Problema", "Valor que aporto"],
    [
        ["Empresas de la UTE", "Versiones distintas, dependencia del sistema de otra empresa y disputas sobre cambios.", "Registro compartido, firma por MSP, historial común y datos privados por lote."],
        ["Administración contratante", "Dificultad para auditar el origen de una certificación y separar ejecución de autorización.", "Trazabilidad de creador/endosantes y participación obligatoria en pagos."],
        ["Dirección de obra / auditoría", "Reconstrucción manual de decisiones y evidencias.", "Explorer con bloque, txId, hashes, función, creador y endosantes."],
        ["Entidad financiera", "Orden de pago desacoplada del estado contractual.", "Evento estructurado PagoAutorizado e integración automática."],
    ],
    [3.5, 6.2, 6.2],
)
add_para(
    doc,
    "La propuesta de valor se resume en cuatro conceptos: confianza compartida, auditoría verificable, confidencialidad selectiva y automatización. La plataforma no pretende que todos vean todo. Pretende que todos vean el hecho común que necesitan coordinar y que solo los socios autorizados accedan al detalle sensible.",
)

doc.add_heading("3.4. Modelo de implantación", level=2)
add_para(
    doc,
    "Plantearía la solución como un servicio B2B/B2G para cada proyecto o UTE. Una entidad neutral —la propia UTE, una ingeniería o un proveedor tecnológico— operaría la infraestructura base, mientras cada organización conservaría sus identidades y su capacidad de validar. El ingreso podría combinar una puesta en marcha por proyecto, una cuota mensual por operación y servicios de integración con ERP, gestor documental y banco.",
)
add_para(
    doc,
    "El producto mínimo viable sería el implementado: canal del proyecto, cinco organizaciones, hitos, pagos, incidencias, evidencias y panel de auditoría. Una segunda fase incorporaría Fabric CA, identidad individual, firma hardware, conectores bancarios y almacenamiento documental regulado. La expansión podría realizarse por nuevas obras, no mediante una única red global: cada proyecto mantendría su gobierno y ciclo de vida.",
)

doc.add_heading("3.5. Beneficios y costes", level=2)
add_bullets(
    doc,
    [
        ("Beneficios operativos. ", "Menos conciliación manual, menos discusiones sobre versiones y seguimiento más rápido del estado de pagos e incidencias."),
        ("Beneficios de control. ", "Separación verificable entre quien ejecuta y quien autoriza, además de una pista de auditoría técnica."),
        ("Beneficios de integración. ", "Eventos que permiten conectar el ledger con sistemas externos sin exponer las claves al frontend."),
        ("Costes. ", "Infraestructura, operación de certificados, monitorización, desarrollo de integraciones y gobierno entre organizaciones."),
        ("Condición de viabilidad. ", "La solución tiene sentido cuando participan entidades realmente independientes; para una sola empresa, una base de datos convencional sería más económica."),
    ],
)

doc.add_heading("3.6. Riesgos del plan", level=2)
add_para(
    doc,
    "El principal riesgo no es tecnológico, sino de gobierno. Las organizaciones deben acordar quién incorpora miembros, cómo se actualizan los chaincodes, qué política exige cada proceso y cómo se resuelven errores. También existen riesgos de protección de datos, custodia de claves y validez jurídica de la evidencia. Fabric aporta trazabilidad técnica, pero no convierte por sí solo un dato falso de origen en verdadero.",
)
add_para(
    doc,
    "Por ese motivo, he concebido la blockchain como una capa de confianza sobre procesos y evidencias, no como sustituto de contratos, controles de obra o normativa de contratación. El valor aparece cuando las reglas técnicas reflejan un acuerdo organizativo previo.",
)

# 4. Arquitectura
add_main_heading(doc, "4. Arquitectura técnica", 4)
doc.add_heading("4.1. Arquitectura general", level=2)
add_para(
    doc,
    "He dividido el sistema en tres planos. El frontend se ejecuta de forma independiente y solo conoce la API HTTPS. La API contiene la autenticación, la orquestación y el Fabric Gateway. La red Fabric mantiene el ledger y ejecuta los contratos. La monitorización se despliega aparte y consume únicamente las métricas de operaciones por la red privada.",
)
add_code(
    doc,
    """
Usuario
  │ HTTPS + JWT
  ▼
Frontend Next.js (Vercel)
  │ REST/JSON
  ▼
API Express + Fabric Gateway + listeners (VM fabric-ute)
  │ gRPC TLS
  ├── peer0 Empresa A ─┐
  ├── peer0 Empresa B  │
  ├── peer0 Empresa C  ├── canal channel-obra ── orderer1/2/3 Raft
  ├── peer0 Empresa D  │
  └── peer0 Administración ┘

GCS almacena evidencias · Prometheus/Grafana observa por VPC
""",
    "Figura 1. Vista lógica de la arquitectura implementada.",
)

doc.add_heading("4.2. Organizaciones, certificados e identidades", level=2)
add_para(
    doc,
    "La red de aplicación contiene cinco organizaciones: EmpresaAMSP, EmpresaBMSP, EmpresaCMSP, EmpresaDMSP y AdministracionMSP. A ellas se añade OrdererMSP para el servicio de ordenación. Cada entidad técnica posee certificados X.509 y claves emitidos por la autoridad de su organización. El MSP no es una base de usuarios: es el conjunto de reglas y raíces de confianza con las que Fabric decide si una identidad pertenece a una organización y qué unidad organizativa representa.",
)
add_para(
    doc,
    "He habilitado NodeOUs para distinguir client, peer, admin y orderer. cryptogen genera un administrador y un usuario de prueba por organización de aplicación, un certificado de identidad para cada peer y certificados de identidad/TLS para los tres orderers. Los certificados TLS aseguran los canales gRPC; los certificados MSP firman propuestas, endosos y operaciones administrativas.",
)
add_table(
    doc,
    ["MSP", "Entidad y papel funcional", "Identidades/nodo"],
    [
        ["EmpresaAMSP", "Constructora A; obra gruesa y cimentación; 35 %.", "Admin, User1 y peer0."],
        ["EmpresaBMSP", "Constructora B; quirófanos e instalaciones; 25 %.", "Admin, User1 y peer0."],
        ["EmpresaCMSP", "Constructora C; socia de A en obra gruesa; 20 %.", "Admin, User1 y peer0."],
        ["EmpresaDMSP", "Constructora D; socia de B en quirófanos; 20 %.", "Admin, User1 y peer0."],
        ["AdministracionMSP", "Ayuntamiento; supervisa y autoriza/rechaza pagos.", "Admin, User1 y peer0."],
        ["OrdererMSP", "Servicio de ordenación.", "Admin y orderer1, orderer2, orderer3."],
    ],
    [3.2, 8.2, 5.0],
)
add_callout(
    doc,
    "Limitación de identidad. ",
    "Aunque cryptogen emite usuarios de prueba, el Gateway actual carga Admin@organización para representar las sesiones de ese MSP. Por tanto, Fabric atribuye correctamente la organización creadora, pero no diferencia criptográficamente a personas dentro de ella. El JWT sí distingue la cuenta en la aplicación. En producción sustituiría este modelo por Fabric CA y certificados individuales protegidos por HSM.",
    LIGHT_GOLD,
)

doc.add_heading("4.3. Peers y orderers", level=2)
add_para(
    doc,
    "He levantado un peer por organización. Cada peer valida y conserva el ledger del canal, mantiene el world state, puede ejecutar chaincode cuando lo tiene instalado, firma endosos y distribuye datos mediante gossip. Los peers de A y Administración forman el perfil diario. Los de B, C y D se añaden en el perfil completo y son imprescindibles cuando una política o colección privada requiere a esas organizaciones.",
)
add_para(
    doc,
    "Los tres orderers pertenecen a OrdererMSP y usan etcdraft. Raft elige un líder y replica el log entre tres votantes; con dos disponibles existe quórum. Su función es establecer un orden total de transacciones, agruparlas en bloques y entregarlas a los peers. No ejecutan chaincode ni deciden la regla de negocio. He mantenido los tres orderers incluso en el perfil diario porque reducirlos a uno eliminaría precisamente la tolerancia que quería demostrar.",
)
add_table(
    doc,
    ["Nodo", "Puerto gRPC", "Puerto de métricas", "Responsabilidad"],
    [
        ["peer0.empresaa", "7051", "9444", "Ledger, endoso de A y PDC obra-gruesa-solar."],
        ["peer0.empresab", "8051", "9446", "Ledger, endoso de B y PDC quirofanos-tech."],
        ["peer0.empresac", "11051", "9447", "Ledger, endoso de C y PDC obra-gruesa-solar."],
        ["peer0.empresad", "12051", "9448", "Ledger, endoso de D y PDC quirofanos-tech."],
        ["peer0.administracion", "9051", "9445", "Ledger y endoso de la Administración; sin PDC de lotes."],
        ["orderer1 / 2 / 3", "7050 / 8050 / 9050", "8443 / 8444 / 8445", "Consenso Raft y creación de bloques."],
    ],
    [4.2, 3.1, 3.5, 7.0],
)

doc.add_heading("4.4. Canal y ledger", level=2)
add_para(
    doc,
    "Toda la aplicación trabaja sobre un único canal, channel-obra, creado con el perfil UteFull. El bloque de configuración incluye las cinco organizaciones aunque en desarrollo solo estén encendidos dos peers. Esta decisión evita mantener una red simplificada con políticas incompatibles y otra distinta para la demostración.",
)
add_para(
    doc,
    "En cada peer, el ledger combina dos estructuras. La blockchain es el historial append-only de bloques y transacciones; conserva la secuencia y los hashes. El world state es una proyección de los valores vigentes que permite consultar sin recorrer todos los bloques. He configurado GoLevelDB como base de estado. No uso CouchDB porque mis consultas se resuelven con rangos, prefijos, claves compuestas y paginación; así reduzco memoria y complejidad.",
)
add_para(
    doc,
    "Los contratos crean claves como hito:<id>, pago:<id>, inc:<id> y estado:obra, además de índices compuestos por estado, empresa, lote o hito. Cuando cambia un hito, Fabric actualiza el world state, pero el valor anterior continúa demostrado por el historial de bloques. Las colecciones privadas mantienen una base privada adicional en los peers autorizados; el canal distribuye el hash de la escritura privada para que los demás puedan verificar su existencia e integridad sin leer el contenido.",
)

doc.add_heading("4.5. Políticas de endoso", level=2)
add_para(
    doc,
    "Las políticas de endoso traducen el acuerdo de gobierno a firmas de peers. No las he codificado dentro de las clases TypeScript; las asigno en el commit de la definición del chaincode. Esto es importante porque permite cambiar el gobierno mediante el lifecycle de Fabric sin recompilar la lógica de negocio.",
)
add_table(
    doc,
    ["Ámbito", "Política", "Razón"],
    [
        ["Lifecycle", "OutOf(2, A, B, C, D, Administración)", "Dos organizaciones bastan para aprobar y comprometer una definición."],
        ["Hito", "OR(A, B, C, D)", "La empresa responsable endosa el avance. Al completar, el mismo tx también debe cumplir Pago."],
        ["Pago", "OR(AND(A,Admin), AND(B,Admin), AND(C,Admin), AND(D,Admin))", "Impide que una empresa libere unilateralmente un pago."],
        ["Incidencia", "OutOf(2, A, B, C, D, Administración)", "Exige dos firmas y se combina con la política de la colección privada."],
        ["Estado de obra", "OR(A, B, C, D, Administración)", "Puede escribirlo la organización que solicita el recálculo."],
        ["PDC obra gruesa", "OR(A, C)", "Solo un socio de ese lote puede endosar la escritura privada."],
        ["PDC quirófanos", "OR(B, D)", "Solo un socio de ese lote puede endosar la escritura privada."],
    ],
    [3.3, 8.2, 6.0],
)
add_code(
    doc,
    """
if [[ "${CC_NAME}" == "pago" ]]; then
  POLICY="OR(
    AND('EmpresaAMSP.peer','AdministracionMSP.peer'),
    AND('EmpresaBMSP.peer','AdministracionMSP.peer'),
    AND('EmpresaCMSP.peer','AdministracionMSP.peer'),
    AND('EmpresaDMSP.peer','AdministracionMSP.peer')
  )"
fi
""",
    "Fragmento de network/scripts/deploy-chaincode.sh. Política de endoso de PagoContract.",
)
add_para(
    doc,
    "La API selecciona expresamente las organizaciones endosantes. Para un pago solicita el par formado por el MSP de la empresa del hito y AdministracionMSP. Para completar un hito usa ese mismo par: la firma de la empresa satisface HitoContract y las dos firmas satisfacen PagoContract, invocado dentro de la misma transacción. No he implementado state-based endorsement por clave; el control se realiza a nivel de definición de chaincode y colección.",
)

doc.add_heading("4.6. Colecciones privadas e incidencias", level=2)
add_para(
    doc,
    "He creado dos Private Data Collections. obra-gruesa-solar pertenece a Empresa A y Empresa C; quirofanos-tech pertenece a Empresa B y Empresa D. memberOnlyRead y memberOnlyWrite están activados, blockToLive es cero —los datos no caducan— y requiredPeerCount es cero para que el perfil diario pueda operar con disponibilidad reducida cuando exista al menos un socio capaz de endosar.",
)
add_para(
    doc,
    "Al abrir una incidencia, el identificador, título, empresa, lote, estado y fechas se escriben en el world state público. El detalle, el coste estimado y las notas técnicas viajan en el transient map, no forman parte de la propuesta pública y se guardan con putPrivateData en la colección del lote. La Administración y el socio del otro lote ven el registro público y el hash privado, pero no obtienen el contenido.",
)
add_code(
    doc,
    """
{
  "name": "obra-gruesa-solar",
  "policy": "OR('EmpresaAMSP.member', 'EmpresaCMSP.member')",
  "memberOnlyRead": true,
  "memberOnlyWrite": true,
  "endorsementPolicy": {
    "signaturePolicy": "OR('EmpresaAMSP.member', 'EmpresaCMSP.member')"
  }
}
""",
    "Fragmento de network/collections-config.json. Colección privada del lote de obra gruesa.",
)

doc.add_heading("4.7. Chaincode de hitos", level=2)
add_para(
    doc,
    "HitoContract gestiona el ciclo PENDIENTE → EN_EJECUCION → VALIDACION → COMPLETADO, con RECHAZADO como salida desde validación. Valida identificador, empresa e importe; impide duplicados; registra fechas deterministas a partir del timestamp de Fabric; y mantiene índices compuestos. Para completar exige una huella SHA-256 de 64 caracteres.",
)
add_para(
    doc,
    "La operación de mayor peso es completarHito. Después de validar la transición, almacena el hash de la evidencia y el txId, actualiza el hito e invoca PagoContract:ponerEnCustodia en el mismo canal y con el mismo contexto transaccional. Si falla cualquiera de las dos escrituras, la transacción completa no se valida.",
)
add_code(
    doc,
    """
async completarHito(ctx: Context, id: string, hashEvidencia: string): Promise<string> {
  const hito = await this.mustGet(ctx, id);
  const hash = this.assertHashEvidencia(hashEvidencia);
  hito.estado = 'COMPLETADO';
  hito.hashEvidencia = hash;
  hito.txId = ctx.stub.getTxID();
  await this.save(ctx, hito, previo);
  const pago = await this.invokeJson(ctx, 'pago', [
    'PagoContract:ponerEnCustodia',
    `pago-${id}`, hito.id, hito.empresa, String(hito.importe), 'completarHito',
  ]);
  return JSON.stringify({ hito, pago });
}
""",
    "Fragmento de chaincode/hito/src/hito-contract.ts. Finalización y custodia en una transacción.",
)

doc.add_heading("4.8. Chaincode de pagos", level=2)
add_para(
    doc,
    "PagoContract inicializa las participaciones de la UTE —35 %, 25 %, 20 % y 20 %— y calcula el desglose monetario de cada pago. ponerEnCustodia impide duplicar el pago o crear dos pagos para un mismo hito. Cuando se invoca de forma independiente, comprueba que el hito esté completado y que coincidan empresa e importe. En el camino hito→pago evita una consulta anidada de vuelta al mismo hito porque Fabric rechazaría la recursión con el mismo txId.",
)
add_para(
    doc,
    "El pago comienza en CUSTODIA y solo puede pasar a AUTORIZADO o RECHAZADO. autorizarPago actualiza el estado y emite PagoAutorizado con pagoId, hitoId, empresa, importe y desglose. El listener del backend escucha el evento y hace POST al banco simulado.",
)
add_code(
    doc,
    """
pago.estado = 'AUTORIZADO';
await ctx.stub.putState(this.key(pagoId), Buffer.from(JSON.stringify(pago)));
ctx.stub.setEvent(
  'PagoAutorizado',
  Buffer.from(JSON.stringify({
    pagoId: pago.id,
    hitoId: pago.hitoId,
    empresa: pago.empresa,
    importeTotal: pago.importeTotal,
    desglose: pago.desglose,
    estado: pago.estado,
  })),
);
""",
    "Fragmento de chaincode/pago/src/pago-contract.ts. Cambio de estado y evento.",
)

doc.add_heading("4.9. Atomicidad y banco simulado", level=2)
add_para(
    doc,
    "Distingo dos límites de atomicidad. Dentro de Fabric, completar el hito y crear el pago en custodia forman una única transacción. No puede quedar el hito COMPLETADO sin su pago CUSTODIA si el invoke del segundo contrato falla. Esta fue una mejora sobre la primera versión, que realizaba dos submit separados desde la API.",
)
add_code(
    doc,
    """
Constructora        API/Fabric Gateway        HitoContract         PagoContract
    │ POST completar        │                      │                     │
    ├──────────────────────►│ submitAsync          │                     │
    │                       ├─────────────────────►│ COMPLETADO          │
    │                       │                      ├──invokeChaincode────►│
    │                       │                      │                     │ CUSTODIA
    │                       │◄──────────── un único commit ──────────────┤

Administración autoriza después → PagoAutorizado → listener API → banco simulado
""",
    "Figura 2. Secuencia del flujo hito–custodia y autorización posterior.",
)
add_para(
    doc,
    "La notificación bancaria, en cambio, es asíncrona y no forma parte del commit de Fabric. El ledger queda como fuente de verdad aunque el webhook falle. El listener reabre el stream tras un error, pero el mock guarda los eventos únicamente en memoria. Para producción incorporaría entrega idempotente, persistencia de eventos pendientes, reintentos con backoff y conciliación bancaria.",
)

doc.add_heading("4.10. Chaincode de incidencias", level=2)
add_para(
    doc,
    "IncidenciaContract valida el lote contra una lista cerrada, comprueba que el MSP cliente pertenezca a la colección y exige el detalle privado en el transient map. El estado público sigue ABIERTA → EN_TRATAMIENTO → CERRADA, con posibilidad de RECHAZADA. La consulta privada repite la comprobación de MSP antes de llamar a getPrivateData.",
)
add_para(
    doc,
    "La API completa esta lógica con reglas de autoría: solo la empresa creadora puede tratar, cerrar, rechazar o adjuntar mientras la incidencia esté abierta o en tratamiento. La colección controla la confidencialidad de red y la API controla la acción de usuario. Son controles complementarios, no equivalentes.",
)

doc.add_heading("4.11. Chaincode de estado de obra", level=2)
add_para(
    doc,
    "EstadoObraContract conserva un documento agregado con totales de hitos, pagos, importes e incidencias. He decidido que el backend lea las listas, calcule el agregado y envíe escribirEstado. Evité llamadas cruzadas desde este contrato porque habrían añadido dependencias, problemas de identidad y timeouts sin mejorar la demostración.",
)
add_para(
    doc,
    "El avance se calcula dividiendo los hitos completados entre el conjunto considerado vigente. Un hito cuyo pago se rechaza se excluye de ese conjunto y se cuenta como rechazado; un hito que ya está en estado RECHAZADO permanece en el total, pero no suma como completado. Esta lógica reside en backend/src/estado.ts y el chaincode valida y persiste números no negativos. La consecuencia es que el ledger demuestra qué agregado se publicó, pero no recalcula por sí solo su procedencia; en una evolución crítica trasladaría o verificaría más reglas dentro de la red.",
)

doc.add_heading("4.12. Evidencias fuera de cadena", level=2)
add_para(
    doc,
    "Guardar fotografías y PDF dentro del ledger multiplicaría el tamaño de cada réplica y empeoraría rendimiento y copias. He optado por almacenar el binario en disco durante el desarrollo y en GCS en producción. La API acepta imágenes o PDF de hasta 5 MB, normaliza el nombre, calcula SHA-256 y conserva metadatos de tamaño, MIME, organización y fecha.",
)
add_para(
    doc,
    "El acta de un hito queda anclada en el campo hashEvidencia del hito al completarlo. En una incidencia creada con fichero, el frontend calcula la huella antes del alta y la incluye en notasTecnicas del PDC; después sube el binario. Esto permite volver a calcular el hash del archivo y compararlo con la huella registrada.",
)
add_callout(
    doc,
    "Límite actual. ",
    "Los adjuntos añadidos a una incidencia después de crearla se registran en el índice off-chain, pero la ruta actual no actualiza el PDC con su hash. Además, el alta de incidencia y la subida del fichero son dos peticiones HTTP: un fallo en la segunda puede dejar el hash anclado sin binario disponible. El anclaje en Fabric está garantizado para el acta de hito y para la evidencia incluida en el alta de incidencia. Una versión definitiva debería unificar el flujo, anclar cada adjunto posterior y almacenar el índice de metadatos en una base persistente, no en index.json local.",
    LIGHT_GOLD,
)

doc.add_heading("4.13. Control de acceso en tres niveles", level=2)
add_para(
    doc,
    "He separado el control de acceso en tres capas. La interfaz oculta acciones incompatibles con el perfil para evitar errores. La API valida el JWT y aplica guardas como perfilConstructora, requireAdministracion, requireEmpresaHito, requireCreadoraIncidencia y requireSocioEvidencia. Finalmente, Fabric valida identidad MSP, endosos y pertenencia a PDC. Una petición que supera la UI aún debe superar la API y la red.",
)
add_table(
    doc,
    ["Regla", "UI/API", "Fabric"],
    [
        ["Avanzar un hito", "Solo una constructora y solo si coincide con la empresa del hito.", "Firma del MSP y política OR de constructoras."],
        ["Autorizar/rechazar pago", "Solo sesión AdministracionMSP.", "Endoso empresa responsable + Administración."],
        ["Tramitar incidencia", "Solo empresa creadora.", "Política 2 de 5; el chaincode no modela la autoría por sí solo."],
        ["Leer detalle privado", "API comprueba socio del lote.", "PDC memberOnlyRead y comprobación de MSP en chaincode."],
        ["Adjuntar evidencia", "Estado y autoría validados; MIME y tamaño limitados.", "El hash se ancla solo en los flujos descritos."],
    ],
    [4.2, 6.4, 7.0],
)
add_para(
    doc,
    "Los roles están parcialmente aplicados en la API y no de forma completa en cada chaincode. Por ejemplo, HitoContract valida que la empresa sea una de las cuatro, pero no comprueba que el MSP cliente corresponda a ese valor; confía en la API y en la selección de endosantes. Además, la ruta de alta acepta un campo empresa opcional y debería forzarlo siempre al perfil autenticado. Un cliente que accediera directamente a Fabric no pasaría por las guardas Express. Lo considero una limitación relevante del prototipo y una prioridad de endurecimiento.",
)

doc.add_heading("4.14. Explorer, inmutabilidad y trazabilidad", level=2)
add_para(
    doc,
    "He implementado un Explorer propio dentro de la API. Un listener comienza en el bloque cero, deserializa cada envelope y extrae número, previousHash, dataHash, txId, timestamp, tipo, chaincode, función, MSP creador y MSP endosantes. El frontend consulta el snapshot cada tres segundos y permite desplegar el detalle y copiar los hashes.",
)
add_para(
    doc,
    "Esta pantalla hace visible la relación entre operación de negocio y evidencia criptográfica. Los orderers no aparecen como endosantes porque su trabajo es ordenar; las firmas mostradas corresponden a peers. Cuando detecto un génesis con un dataHash distinto, vacío el snapshot para no mezclar dos redes creadas tras un reset. El índice se mantiene en memoria, por lo que no sustituye a Hyperledger Explorer ni a una base de indexación persistente, pero resulta suficiente para demostrar la trazabilidad.",
)
add_para(
    doc,
    "No existe una operación de borrado de hitos ni un botón para reiniciar el avance. Reutilizar un identificador produce error. Empezar de cero implica crear una red nueva y borrar sus volúmenes, no editar la historia. Esta diferencia entre reinicializar un entorno de demostración y modificar el ledger es esencial para explicar la inmutabilidad con rigor.",
)

doc.add_heading("4.15. Despliegue y seguridad de red", level=2)
add_para(
    doc,
    "En producción coloqué la red y la API en fabric-ute. Caddy termina TLS con un certificado de Let's Encrypt y reenvía hacia ute-api:4000. Los certificados de peer incluyen nombre DNS, FQDN e IP pública como SAN. La API permite el origen del frontend Vercel mediante CORS y publica health, Swagger y las rutas de negocio por HTTPS.",
)
add_para(
    doc,
    "Los puertos gRPC de peers y orderers no se exponen a Internet. Los puertos de operaciones se limitan por firewall a 10.8.0.0/24, desde donde la VM monitoring-ute los recopila. GCS usa la cuenta de servicio de Compute Engine y el repositorio excluye .env, certificados, claves, volúmenes y evidencias. Estas medidas no constituyen una auditoría de seguridad completa, pero reducen la superficie del prototipo.",
)

doc.add_heading("4.16. Monitorización y alertas", level=2)
add_para(
    doc,
    "Prometheus consulta cada 15 segundos las métricas de peers, orderers y backend. Grafana dispone de un dashboard provisionado y se integra en la ruta /monitor. Definí tres reglas alineadas con los riesgos principales: PeerCaido tras un minuto sin respuesta, LatenciaBloqueAlta cuando el p99 supera cinco segundos y ErrorEndorsementAlto cuando más del 5 % de las propuestas falla durante cinco minutos.",
)
add_para(
    doc,
    "La alerta de peer caído se validó deteniendo realmente peer0 de Empresa A. En una VM e2-standard-4, la latencia y el error reales no alcanzaban de forma natural los umbrales, por lo que usé un exporter de demostración para inducir ambas series. No presento estos dos disparos como incidentes reales: demuestran que las reglas y el panel funcionan cuando reciben datos por encima del umbral.",
)

doc.add_heading("4.17. Pruebas y consumo", level=2)
add_para(
    doc,
    "La suite unitaria cubre creación, duplicados, importes, transiciones, hashes, custodia, desglose, autorización, evento, rechazo, PDC, permisos MSP, transient data, paginación y persistencia del agregado. La ejecución final fue satisfactoria en los cuatro paquetes: 43 pruebas y cuatro suites.",
)
add_para(
    doc,
    "Complementé Jest con verificaciones de canal y nodos, ciclo hito-pago, API autenticada, PDC, frontend y seed completo. El flujo de nube produjo hitos completados, pagos en custodia y autorizados, eventos en el banco simulado y bloques visibles en el Explorer.",
)
add_para(
    doc,
    "No existen suites automatizadas específicas para el backend o el frontend. Los scripts ejercitan la API y comprueban que la UI responda por HTTP, pero no automatizan la interacción de un navegador con Playwright, Cypress o Selenium. Por tanto, describo esta cobertura como integración y smoke test, no como una prueba end-to-end visual completa.",
)
add_para(
    doc,
    "También medí memoria en vez de limitarme a estimarla. Instalar los cuatro chaincodes en los cinco peers crea veinte contenedores fabric-nodeenv y eleva el conjunto a unos 1,7 GB. Instalando chaincode solo en los peers que endosan, el escenario medido utilizó diez nodeenv y aproximadamente 1.024 MiB de contenedores. Esta optimización hizo viable la demostración en WSL2 y mostró la importancia del diseño operativo.",
)

doc.add_heading("4.18. Decisiones, desviaciones y limitaciones", level=2)
add_table(
    doc,
    ["Decisión", "Justificación", "Consecuencia"],
    [
        ["VM GCP en vez de Railway/Render", "El listener de eventos debe permanecer activo y cerca de los peers; los planes gratuitos pueden suspender procesos.", "Mayor control y coste gestionado apagando la VM fuera de ensayos."],
        ["cryptogen en vez de Fabric CA", "Plazo y alcance académico.", "Identidades estáticas de prueba; sin alta/revocación real."],
        ["LevelDB en vez de CouchDB", "Consultas por claves compuestas y menor consumo.", "Sin rich queries JSON."],
        ["Estado calculado en backend", "Evitar dependencias y timeouts entre varios chaincodes.", "El agregado se registra, pero el contrato no verifica toda su derivación."],
        ["Evidencia off-chain", "No replicar megabytes en cada peer.", "Hay que gobernar almacenamiento, metadatos y disponibilidad fuera de Fabric."],
        ["Explorer en memoria", "Demostración directa y ligera.", "Reconstruye desde el bloque 0 al reiniciar y no es un producto de exploración completo."],
    ],
    [4.0, 7.0, 6.5],
)
add_para(
    doc,
    "Las limitaciones principales son las identidades de prueba, el banco simulado, la persistencia en memoria del webhook, los roles parcialmente trasladados a chaincode, el índice local de evidencias, la integración Vercel pendiente de un redeploy confirmado y la falta de pruebas de carga. La alta disponibilidad se demuestra en el consenso de tres orderers, pero no en peers redundantes, API o almacenamiento de metadatos.",
)

# 5. Manual de usuario
add_main_heading(doc, "5. Manual de usuario", 5)
doc.add_heading("5.1. Acceso y perfiles", level=2)
add_para(
    doc,
    "Al abrir la aplicación encuentro la pantalla de acceso institucional. Puedo elegir cualquiera de las cinco cuentas de demostración; para facilitar la defensa, la contraseña coincide con el nombre de usuario. Tras autenticarme, la cabecera muestra la organización, el oficio, el lote y, en las constructoras, el porcentaje de participación. La sesión dura ocho horas y el botón Salir elimina el token.",
)
add_table(
    doc,
    ["Usuario", "Perfil", "Acciones principales"],
    [
        ["empresaA", "Empresa A · obra gruesa · 35 %", "Hitos propios, incidencias del lote A/C y lectura de su PDC."],
        ["empresaB", "Empresa B · quirófanos · 25 %", "Hitos propios, incidencias del lote B/D y lectura de su PDC."],
        ["empresaC", "Empresa C · obra gruesa · 20 %", "Hitos propios, incidencias del lote A/C y lectura de su PDC."],
        ["empresaD", "Empresa D · quirófanos · 20 %", "Hitos propios, incidencias del lote B/D y lectura de su PDC."],
        ["administracion", "Ayuntamiento", "Consulta pública y autorización o rechazo de pagos."],
    ],
    [3.2, 6.2, 8.0],
)

doc.add_heading("5.2. Navegación", level=2)
add_para(
    doc,
    "El menú superior da acceso a Inicio, Hitos, Pagos, Incidencias, Estado obra, Explorer y Monitor. La aplicación comprueba que exista sesión antes de mostrar estas páginas. Los mensajes de error separan un rechazo esperado de negocio —por ejemplo, intentar abrir un PDC ajeno— de una indisponibilidad de infraestructura, y conservan el detalle técnico desplegable.",
)

doc.add_heading("5.3. Pantalla Inicio", level=2)
add_para(
    doc,
    "En Inicio consulto un resumen del proyecto: porcentaje de avance, número de hitos, pagos en custodia e incidencias abiertas. Debajo veo cuánto importe autorizado corresponde a cada empresa según las participaciones y el total de pagos autorizados. Los accesos directos me llevan a Hitos o al Explorer.",
)

doc.add_heading("5.4. Pantalla Hitos", level=2)
add_numbered(
    doc,
    [
        "Entro con una cuenta de constructora. Administración solo ve la información pública y no dispone de formulario de alta.",
        "Indico título e importe. La aplicación asigna la empresa del perfil y crea el hito en PENDIENTE.",
        "Pulso Iniciar para pasar a EN_EJECUCION y Validar para llegar a VALIDACION. Solo aparece la acción válida para el estado actual.",
        "En VALIDACION selecciono una fotografía o PDF de hasta 5 MB. El botón Completar permanece deshabilitado mientras no haya archivo.",
        "Al completar, el fichero queda fuera de Fabric, su SHA-256 se guarda en el hito y se crea pago-<id> en CUSTODIA dentro de la misma transacción.",
        "La ficha muestra estado, empresa, importe, pago asociado, txId, número de bloque, hash del acta y opción de descargar la evidencia.",
    ],
)
add_para(
    doc,
    "A la derecha aparece un Explorer compacto con los últimos ocho bloques. Esto me permite comprobar de inmediato que la acción de la pantalla ha generado actividad en el canal.",
)

doc.add_heading("5.5. Pantalla Pagos", level=2)
add_para(
    doc,
    "La lista muestra cada pago, su hito, empresa, importe total, estado y desglose 35/25/20/20. Si entro como constructora y el pago está en CUSTODIA, leo que espera autorización y no tengo botones. Si entro como Administración aparecen Autorizar y Rechazar.",
)
add_para(
    doc,
    "Al autorizar, la pantalla confirma la transición CUSTODIA → AUTORIZADO y explica el evento emitido. Para comprobar la integración puedo consultar GET /mock/banco/pagos. Al rechazar, el estado pasa a RECHAZADO y no se emite PagoAutorizado.",
)

doc.add_heading("5.6. Pantalla Incidencias", level=2)
add_numbered(
    doc,
    [
        "Entro como constructora. El formulario informa de la empresa y del lote que corresponden a mi sesión.",
        "Indico título, detalle técnico, coste estimado y, opcionalmente, una fotografía o PDF.",
        "La aplicación registra en el canal la parte pública y envía detalle, coste y notas mediante transient data a la colección privada.",
        "Pulso Ver PDC. Si soy socio del lote, veo detalle, coste y hash; si no lo soy, recibo una denegación que demuestra el aislamiento.",
        "Solo la empresa que abrió la incidencia puede usar en la interfaz Tratar y después Cerrar. La API también implementa Rechazar, aunque la pantalla actual no muestra ese botón.",
        "Mientras esté abierta o en tratamiento, la creadora puede adjuntar evidencias. Los socios del lote pueden listar y descargar las autorizadas por la API.",
    ],
)
add_para(
    doc,
    "La pantalla consulta /red para detectar qué peers están activos. Si ningún socio del lote tiene peer disponible, muestra el comando make pdc-up antes de que el usuario intente una operación imposible. En el perfil diario, el lote de quirófanos necesita peer B o D.",
)

doc.add_heading("5.7. Pantalla Estado de obra", level=2)
add_para(
    doc,
    "Esta vista reúne avance, hitos completados y rechazados, pagos en custodia y autorizados, importes e incidencias. El indicador de pagos compara lo autorizado con el importe previsto de los hitos vigentes. Puedo pulsar Recalcular para que el backend vuelva a leer las listas y escriba un nuevo agregado en EstadoObraContract.",
)

doc.add_heading("5.8. Pantalla Explorer", level=2)
add_para(
    doc,
    "Explorer muestra la altura del canal y pagina los bloques de veinte en veinte. Al desplegar un bloque veo dataHash, previousHash y sus transacciones. Cada transacción muestra txId, hora, chaincode, función, MSP creador y MSP endosantes. Los valores pueden copiarse para comparar una operación de negocio con la evidencia de Fabric.",
)

doc.add_heading("5.9. Pantalla Monitor", level=2)
add_para(
    doc,
    "Monitor integra el dashboard de Grafana en un iframe. Desde él observo disponibilidad de peers, latencia de procesamiento de bloques, propuestas de endoso y tabla de alertas. En producción, la URL se sirve bajo el mismo dominio HTTPS mediante el proxy de Caddy.",
)

doc.add_heading("5.10. Flujo recomendado de demostración", level=2)
add_numbered(
    doc,
    [
        "Entrar como empresaA y comprobar el perfil de obra gruesa.",
        "Crear un hito, iniciarlo, enviarlo a validación y completarlo con un acta.",
        "Comprobar en la ficha el hash, txId, bloque y pago en custodia; abrir Explorer y localizar completarHito.",
        "Salir, entrar como administracion y autorizar el pago.",
        "Comprobar PagoAutorizado en el banco simulado y revisar los dos MSP endosantes en Explorer.",
        "Entrar como empresaA, crear una incidencia con evidencia y leer el PDC.",
        "Intentar leer el mismo detalle como Administración para mostrar el rechazo esperado.",
        "Con los peers completos, repetir con empresaB en quirofanos-tech y contrastar el acceso desde empresaA.",
        "Recalcular el estado de obra y abrir Monitor para revisar métricas y alertas.",
    ],
)

doc.add_heading("5.11. Operación local", level=2)
add_code(
    doc,
    """
make up-dev       # 3 orderers + peer A + peer Administración + canal
make deploy-cc    # desplegar e inicializar los cuatro chaincodes
make api-up       # API Express
make ui-up        # frontend Next.js
make pdc-up       # añadir peers B/C/D cuando se necesiten PDC completos
make monitoring-up
""",
    "Comandos principales de arranque local.",
)
add_para(
    doc,
    "Para una demostración completa uso make up-full. Si necesito una cadena nueva empleo los objetivos reset-demo, que detienen la API, eliminan volúmenes y evidencias, crean otro ledger y vuelven a desplegar. No utilizo ese procedimiento para alterar datos de una cadena existente.",
)

# 6. Conclusiones
add_main_heading(doc, "6. Conclusiones personales", 6)
doc.add_heading("6.1. Valoración del resultado", level=2)
add_para(
    doc,
    "Me he quedado satisfecho con el resultado porque el proyecto no se limita a mostrar una red Fabric encendida. He conseguido recorrer un proceso reconocible de principio a fin: una empresa registra trabajo, aporta evidencia, completa un hito, genera custodia en la misma transacción, la Administración autoriza con su firma y un sistema externo recibe el evento. Después puedo abrir el Explorer y explicar quién creó la transacción, qué peers la endosaron y en qué bloque quedó.",
)
add_para(
    doc,
    "También valoro especialmente la privacidad selectiva. La solución demuestra que compartir un ledger no obliga a compartir todos los datos. El resto de miembros puede conocer que existe una incidencia y verificar el hash de su información privada, mientras el coste y las notas quedan en los peers de los socios del lote. Esta combinación refleja mejor una colaboración empresarial real que una transparencia absoluta.",
)

doc.add_heading("6.2. Aprendizajes demostrados", level=2)
add_para(
    doc,
    "A lo largo del desarrollo he aplicado de forma conjunta conceptos que en clase suelen estudiarse por separado: PKI y certificados X.509, MSP, TLS, peers, orderers, consenso Raft, lifecycle de chaincode, políticas de endoso, transient data, private data collections, world state, eventos, Gateway y observabilidad. El proyecto me ha obligado a entender no solo qué hace cada pieza, sino cómo una decisión en una capa condiciona a las demás.",
)
add_para(
    doc,
    "La política de pagos ha sido el ejemplo más claro. No bastaba con ocultar un botón a la constructora: tuve que seleccionar dos organizaciones endosantes, dirigir la propuesta al peer adecuado y mantener la política OR de pares empresa–Administración. Del mismo modo, para las incidencias tuve que coordinar política general 2 de 5, política de colección, MSP del cliente y disponibilidad real de peers.",
)
add_para(
    doc,
    "También he aprendido a medir y justificar. Elegí LevelDB por el patrón de consultas y el consumo; instalé chaincode solo donde podía endosar; separé la monitorización porque la VM pequeña era suficiente; y preferí una VM a un servicio que podía dormir el listener. Estas decisiones no son adornos de documentación: explican por qué la solución pudo ejecutarse de forma estable.",
)

doc.add_heading("6.3. Lectura crítica", level=2)
add_para(
    doc,
    "No considero que el prototipo sea todavía una plataforma de producción. La identidad de usuario termina representada por un certificado administrador de su organización, el banco es simulado, algunos permisos dependen de la API, el índice de evidencias sigue siendo local, el Explorer es volátil y las alertas de latencia y endoso se indujeron para la demostración. Además, el último estado documentado de Vercel necesita un redeploy correcto.",
)
add_para(
    doc,
    "Precisamente esta lectura crítica forma parte de lo aprendido. Fabric protege lo que sus políticas expresan; no corrige automáticamente una autorización incompleta. Un hash demuestra integridad desde el anclaje, pero no disponibilidad del archivo. Un evento confirma un cambio de ledger, pero no garantiza por sí solo que un banco haya ejecutado la transferencia. Entender estos límites me parece tan importante como hacer funcionar la demo.",
)

doc.add_heading("6.4. Trabajo futuro", level=2)
add_bullets(
    doc,
    [
        ("Identidad. ", "Sustituir cryptogen por Fabric CA, emitir certificados individuales, revocar credenciales y proteger claves con HSM o KMS."),
        ("Autorización. ", "Trasladar al chaincode las comprobaciones críticas de MSP y atributo; impedir que el body elija una empresa distinta del perfil."),
        ("Banca. ", "Integrar una API bancaria real con OAuth/mTLS, idempotencia, outbox persistente, reintentos y conciliación."),
        ("Documentos. ", "Usar almacenamiento documental definitivo con versionado, retención, metadatos persistentes y anclaje on-chain de todos los adjuntos."),
        ("Alta disponibilidad. ", "Desplegar peers redundantes, API replicada, balanceo, backup y recuperación probada."),
        ("Rendimiento. ", "Ejecutar pruebas de carga, medir throughput/latencia por política y ajustar tamaño de bloque y recursos."),
        ("Operación. ", "Confirmar el despliegue Vercel, endurecer secretos, revisar CORS/rate limits y completar alertas sobre integraciones externas."),
    ],
)

doc.add_heading("6.5. Conclusión final", level=2)
add_para(
    doc,
    "Mi conclusión es que Hyperledger Fabric encaja en este caso no por ser una tecnología novedosa, sino porque resuelve una necesidad concreta de gobierno: varias organizaciones pueden compartir una verdad operativa sin que una de ellas tenga capacidad unilateral para reescribirla. Sobre esa base he construido auditoría, confidencialidad selectiva y automatización del pago.",
)
add_para(
    doc,
    "El resultado final demuestra lo que quería aprender: diseñar una red permisionada desde sus identidades y políticas, conectarla con una aplicación completa, desplegarla fuera del portátil, observarla y defender sus decisiones con evidencias. La plataforma no elimina la necesidad de confianza entre personas y organizaciones, pero convierte parte de esa confianza en reglas técnicas comunes, verificables y trazables.",
)

# Propiedades y guardado
doc.core_properties.title = "Memoria final — Plataforma blockchain para la gestión de una UTE"
doc.core_properties.subject = "Hyperledger Fabric, hitos, pagos, incidencias y estado de obra"
doc.core_properties.author = "Proyecto UTE Blockchain Solutions"
doc.core_properties.keywords = "Hyperledger Fabric, UTE, blockchain, escrow, PDC, GCP"
doc.core_properties.comments = "Generado a partir del código y la documentación técnica del proyecto."

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
